from datetime import datetime, timezone, timedelta
import json

"""Independent Tournament module routes.

V1.4.27 introduces the first real Tournament database core:
- tournaments
- tournament_registrations
- tournament_members

Tournament data stays independent from Rank / Season / normal matches.
"""

TOURNAMENT_AREA_SETTING_KEY = "tournament_area_enabled"
TOURNAMENT_DESIGN_SETTING_KEY = "tournament_design_v1"

DEFAULT_TOURNAMENT_DESIGN = {
    "hero_cup_width": 220,
    "hero_cup_right": 24,
    "hero_cup_bottom": -14,
    "arena_badge_width": 230,
    "arena_badge_x": 50,
    "arena_badge_y": 46,
}


def register_routes(context):
    globals().update(context)

    def tournament_area_enabled(force=False):
        request_key = "_tournament_area_enabled_cached"
        if not force:
            cached = cache_get(request_key)
            if isinstance(cached, bool):
                return cached
            cached = ttl_cache_get("tournament_area_enabled")
            if isinstance(cached, bool):
                return cache_set(request_key, cached)

        enabled = False
        try:
            result = execute_query(
                db.table("system_settings").select("setting_value")
                .eq("setting_key", TOURNAMENT_AREA_SETTING_KEY).limit(1),
                "get_tournament_area_enabled", attempts=2,
            )
            raw = ((result.data or [{}])[0]).get("setting_value")
            if isinstance(raw, dict):
                raw = raw.get("enabled")
            if isinstance(raw, str):
                raw = raw.strip().lower() in {"1", "true", "yes", "on", "enabled"}
            if isinstance(raw, bool):
                enabled = raw
        except Exception as exc:
            print(f"tournament_area_enabled warning: {exc}")

        ttl_cache_set("tournament_area_enabled", enabled, 45)
        return cache_set(request_key, enabled)


    def tournament_design_settings(force=False):
        request_key = "_tournament_design_settings_cached"
        if not force:
            cached = cache_get(request_key)
            if isinstance(cached, dict):
                return cached
            cached = ttl_cache_get("tournament_design_settings")
            if isinstance(cached, dict):
                return cache_set(request_key, cached)

        value = dict(DEFAULT_TOURNAMENT_DESIGN)
        try:
            result = execute_query(
                db.table("system_settings").select("setting_value")
                .eq("setting_key", TOURNAMENT_DESIGN_SETTING_KEY).limit(1),
                "get_tournament_design_settings", attempts=2,
            )
            raw = ((result.data or [{}])[0]).get("setting_value")
            if isinstance(raw, dict):
                for key in value:
                    if key in raw:
                        try:
                            value[key] = int(raw[key])
                        except (TypeError, ValueError):
                            pass
        except Exception as exc:
            app.logger.warning("Tournament design settings unavailable: %s", exc)

        ttl_cache_set("tournament_design_settings", value, 45)
        return cache_set(request_key, value)

    def _safe_rows(query, label):
        try:
            result = execute_query(query, label, attempts=2)
            return [dict(row) for row in (result.data or [])], None
        except Exception as exc:
            app.logger.warning("Tournament DB unavailable [%s]: %s", label, exc)
            return [], str(exc)

    def _list_tournaments():
        rows, error = _safe_rows(
            db.table("tournaments").select("*").eq("is_visible", True).order("created_at"),
            "tournament_list",
        )
        return rows, error

    def _parse_tournament_dt(raw):
        if not raw:
            return None
        try:
            return datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
        except Exception:
            return None

    def _registration_for_user(tournament_id, user_id):
        if not tournament_id or not user_id:
            return None
        rows, _ = _safe_rows(
            db.table("tournament_registrations").select("*")
            .eq("tournament_id", tournament_id).eq("user_id", user_id).limit(1),
            "tournament_registration_for_user",
        )
        return rows[0] if rows else None

    def _tournament_timing_preview(tournament_id):
        if not tournament_id:
            return {}
        rows, _ = _safe_rows(
            db.table("tournament_settings").select("setting_value")
            .eq("tournament_id", tournament_id).eq("setting_key", "competition_timing").limit(1),
            "tournament_timing_preview",
        )
        value = (rows[0].get("setting_value") if rows else {}) or {}
        return value if isinstance(value, dict) else {}

    def _member_for_user(tournament_id, user_id):
        if not tournament_id or not user_id:
            return None
        rows, _ = _safe_rows(
            db.table("tournament_members").select("*")
            .eq("tournament_id", tournament_id).eq("user_id", user_id).limit(1),
            "tournament_member_for_user",
        )
        return rows[0] if rows else None

    def _decorate_registration_rows(rows):
        user_ids = [str(row.get("user_id")) for row in rows if row.get("user_id")]
        users = {}
        if user_ids:
            user_rows, _ = _safe_rows(
                db.table("users").select("id,username,display_name,avatar_url,account_status")
                .in_("id", user_ids),
                "tournament_registration_users",
            )
            users = {str(row.get("id")): row for row in user_rows}
        for row in rows:
            user = users.get(str(row.get("user_id"))) or {}
            row["user"] = user
            row["display_name"] = user.get("display_name") or user.get("username") or "Tài khoản"
        return rows

    def _decorate_finance(row):
        fee_amount = int(row.get("fee_amount") or 50000)
        responsibility_amount = int(row.get("responsibility_amount") or 50000)
        paid = int(row.get("amount_paid") or 0)
        deducted = max(0, int(row.get("responsibility_deducted") or 0))
        refunded = max(0, int(row.get("amount_refunded") or 0))
        required = fee_amount + responsibility_amount
        row["fee_amount"] = fee_amount
        row["responsibility_amount"] = responsibility_amount
        row["amount_paid"] = paid
        row["responsibility_deducted"] = min(deducted, responsibility_amount)
        row["amount_refunded"] = refunded
        row["required_amount"] = required
        row["amount_missing"] = max(0, required - paid)
        row["amount_surplus"] = max(0, paid - required - refunded)
        row["responsibility_refundable"] = max(0, responsibility_amount - row["responsibility_deducted"] - refunded)
        if paid < required:
            row["finance_status"] = "missing"
        elif paid > required + refunded:
            row["finance_status"] = "surplus"
        else:
            row["finance_status"] = "paid"
        return row

    def _admin_tournament_data():
        data = {
            "db_ready": True,
            "tournaments": [],
            "selected": None,
            "pending": [],
            "approved": [],
            "rejected": [],
            "members": [],
        }
        tournaments, error = _list_tournaments()
        if error:
            data["db_ready"] = False
            return data
        data["tournaments"] = tournaments
        selected = tournaments[0] if tournaments else None
        data["selected"] = selected
        if not selected:
            return data
        tournament_id = selected.get("id")
        registrations, _ = _safe_rows(
            db.table("tournament_registrations").select("*")
            .eq("tournament_id", tournament_id).order("registered_at"),
            "admin_tournament_registrations",
        )
        registrations = [_decorate_finance(row) for row in _decorate_registration_rows(registrations)]
        unmatched, _ = _safe_rows(
            db.table("tournament_fee_unmatched").select("*")
            .eq("tournament_id", tournament_id).eq("status", "waiting").order("created_at", desc=True),
            "admin_tournament_fee_unmatched",
        )
        data["fee_unmatched"] = unmatched
        data["pending"] = [row for row in registrations if row.get("status") == "pending"]
        data["approved"] = [row for row in registrations if row.get("status") == "approved"]
        data["rejected"] = [row for row in registrations if row.get("status") == "rejected"]

        finance_rows = data["pending"] + data["approved"]
        unmatched_finance = []
        for fee in unmatched:
            paid = int(fee.get("amount_paid") or 0)
            unmatched_finance.append({
                "display_name": fee.get("payer_name") or "HLV chưa ghép tài khoản",
                "amount_paid": paid,
                "amount_missing": max(0, 100000 - paid),
                "amount_surplus": max(0, paid - 100000),
                "amount_refunded": 0,
                "fee_amount": 50000,
                "responsibility_amount": 50000,
                "is_unmatched_fee": True,
            })
        all_finance_rows = finance_rows + unmatched_finance
        data["finance_summary"] = {
            "player_count": len(all_finance_rows),
            "total_collected": sum(int(r.get("amount_paid") or 0) for r in all_finance_rows),
            "tournament_fee_collected": sum(min(int(r.get("amount_paid") or 0), int(r.get("fee_amount") or 50000)) for r in all_finance_rows),
            "responsibility_collected": sum(max(0, min(int(r.get("amount_paid") or 0) - int(r.get("fee_amount") or 50000), int(r.get("responsibility_amount") or 50000))) for r in all_finance_rows),
            "total_missing": sum(int(r.get("amount_missing") or 0) for r in all_finance_rows),
            "total_surplus": sum(int(r.get("amount_surplus") or 0) for r in all_finance_rows),
            "total_refunded": sum(int(r.get("amount_refunded") or 0) for r in all_finance_rows),
            "surplus_players": [r for r in all_finance_rows if int(r.get("amount_surplus") or 0) > 0],
        }

        members, _ = _safe_rows(
            db.table("tournament_members").select("*")
            .eq("tournament_id", tournament_id).order("approved_at"),
            "admin_tournament_members",
        )
        data["members"] = _decorate_registration_rows(members)
        return data

    @app.context_processor
    def inject_tournament_context():
        payload = {"tournament_area_enabled": tournament_area_enabled(), "tournament_design": tournament_design_settings()}
        if request.endpoint == "admin":
            payload["tournament_admin_data"] = _admin_tournament_data()
        return payload

    @app.get('/dang-ky-c1')
    @login_required
    def tournament_register_shortcut():
        """Short public-friendly link that opens the active C1 registration form."""
        if not tournament_area_enabled():
            flash("Khu vực Giải đấu đang tạm đóng.", "warning")
            return redirect(url_for("tournaments"))

        tournament_rows, error = _list_tournaments()
        if error:
            flash("Chưa thể truy cập dữ liệu giải đấu lúc này.", "error")
            return redirect(url_for("tournaments"))

        active = next((
            item for item in tournament_rows
            if item.get("registration_open") and item.get("status") in {"registration", "upcoming"}
        ), None)
        if not active:
            flash("Hiện chưa có giải đấu nào đang mở đăng ký.", "warning")
            return redirect(url_for("tournaments"))

        tournament_id = active.get("id")
        return redirect(url_for("tournaments", register=tournament_id) + f"#register-{tournament_id}")



    # V1.5.54 - Dữ liệu cá nhân C1 được dựng ngay tại /tournaments để phần
    # Đối thủ / Giờ rảnh / Lịch đối thủ nằm đúng ở trang danh sách giải.
    def _landing_setting(tournament_id, key, default=None):
        rows, _ = _safe_rows(
            db.table("tournament_settings").select("setting_value")
            .eq("tournament_id", tournament_id).eq("setting_key", key),
            f"tournament_landing_setting_{key}",
        )
        return (rows[0] or {}).get("setting_value", default) if rows else default

    def _landing_availability_days():
        vn_tz = timezone(timedelta(hours=7))
        now = datetime.now(vn_tz)
        labels = ("Hôm nay", "Ngày mai", "Ngày kia")
        weekdays = ("Thứ Hai","Thứ Ba","Thứ Tư","Thứ Năm","Thứ Sáu","Thứ Bảy","Chủ nhật")
        out=[]
        for offset, label in enumerate(labels):
            d=now.date()+timedelta(days=offset)
            hours=[11,12,18,19,20,21] if d.weekday()<5 else list(range(11,22))
            slots=[]
            for hour in hours:
                dt=datetime(d.year,d.month,d.day,hour,0,tzinfo=vn_tz)
                if dt>now:
                    slots.append({"iso":dt.isoformat(),"time":f"{hour:02d}:00","label":f"{hour:02d}:00 – {hour+1:02d}:00"})
            out.append({"date":d.isoformat(),"label":label,"weekday":f"{weekdays[d.weekday()]} · {d.strftime('%d/%m')}","is_weekend":d.weekday()>=5,"slots":slots})
        return out

    def _landing_parse_slots(values):
        vn_tz=timezone(timedelta(hours=7)); now=datetime.now(vn_tz); max_day=now.date()+timedelta(days=2)
        out=[]
        for raw in values or []:
            try:
                dt=datetime.fromisoformat(str(raw).replace('Z','+00:00'))
                if dt.tzinfo is None: dt=dt.replace(tzinfo=vn_tz)
                dt=dt.astimezone(vn_tz)
                if dt>now and dt.date()<=max_day:
                    out.append(dt.isoformat())
            except Exception: pass
        return sorted(set(out))

    def _landing_group_slots(values, mine_set):
        days=_landing_availability_days(); grouped={d['date']:[] for d in days}; mine_set=set(mine_set or [])
        for iso in values or []:
            try:
                dt=datetime.fromisoformat(str(iso)); day=dt.date().isoformat()
                if day in grouped:
                    grouped[day].append({"iso":iso,"label":dt.strftime('%H:%M'),"is_overlap":iso in mine_set})
            except Exception: pass
        return [{**d,"opponent_slots":sorted(grouped[d['date']],key=lambda x:x['iso'])} for d in days]

    def _landing_hub_payload(tournament_id, user, member):
        uid=str((user or {}).get('id') or '')
        test_state=_landing_setting(tournament_id,'c1_test_accounts_v1',{}) or {}
        test_ids=[str(x) for x in (test_state.get('user_ids') or [])][:2]
        is_test=uid in set(test_ids)
        if not member and not is_test:
            return None
        days=_landing_availability_days()
        if is_test:
            av_state=_landing_setting(tournament_id,f'c1_test_availability_{uid}',{}) or {}
            mine=_landing_parse_slots(av_state.get('slots') or [])
        else:
            av_rows,_=_safe_rows(db.table('tournament_availability_slots').select('slot_at').eq('tournament_id',tournament_id).eq('user_id',uid),'tournament_landing_mine_availability')
            mine=_landing_parse_slots([r.get('slot_at') for r in av_rows])
        mine_set=set(mine)

        match_rows,_=_safe_rows(db.table('tournament_matches').select('*').eq('tournament_id',tournament_id).or_(f'home_user_id.eq.{uid},away_user_id.eq.{uid}').order('created_at'),'tournament_landing_my_matches') if member else ([],None)
        ids={uid}
        for m in match_rows:
            if m.get('home_user_id'): ids.add(str(m.get('home_user_id')))
            if m.get('away_user_id'): ids.add(str(m.get('away_user_id')))
        if is_test: ids.update(test_ids)
        user_rows,_=_safe_rows(db.table('users').select('id,username,display_name').in_('id',list(ids)),'tournament_landing_users') if ids else ([],None)
        names={str(x.get('id')):(x.get('display_name') or x.get('username') or 'HLV') for x in user_rows}
        reg_rows,_=_safe_rows(db.table('tournament_registrations').select('user_id,zalo_name').eq('tournament_id',tournament_id).in_('user_id',list(ids)),'tournament_landing_zalo') if ids else ([],None)
        zalos={str(x.get('user_id')):(x.get('zalo_name') or '') for x in reg_rows}

        decorated=[]
        for m in match_rows:
            row=dict(m); h=str(row.get('home_user_id') or ''); a=str(row.get('away_user_id') or '')
            row['home_name']=names.get(h,'HLV'); row['away_name']=names.get(a,'HLV'); row['home_zalo_name']=zalos.get(h,''); row['away_zalo_name']=zalos.get(a,'')
            oid=a if h==uid else h
            opp_rows,_=_safe_rows(db.table('tournament_availability_slots').select('slot_at').eq('tournament_id',tournament_id).eq('user_id',oid),'tournament_landing_opp_availability')
            opp=_landing_parse_slots([r.get('slot_at') for r in opp_rows])
            row['opponent_availability_days']=_landing_group_slots(opp,mine_set)
            decorated.append(row)

        test_opponent=None; test_opp_days=[]
        if is_test:
            oid=next((x for x in test_ids if x!=uid),'')
            test_opponent={"id":oid,"display_name":names.get(oid,'HLV')} if oid else None
            if oid:
                state=_landing_setting(tournament_id,f'c1_test_availability_{oid}',{}) or {}
                opp=_landing_parse_slots(state.get('slots') or [])
                test_opp_days=_landing_group_slots(opp,mine_set)

        s1_reveals=_landing_setting(tournament_id,'stage1_player_reveals',{}) or {}
        league_draw=_landing_setting(tournament_id,'league_draw_v2',{}) or {}
        league_reveals=_landing_setting(tournament_id,'league_player_reveals',{}) or {}
        league_mine=(league_draw.get('revealed') or {}).get(uid,[]) if isinstance(league_draw,dict) else []

        # V1.5.54: dữ liệu Trung tâm C1 được đưa ra /tournaments.
        # Trang landing giờ là nơi theo dõi tiến trình, Host, phòng đang chạy và trận của HLV.
        timing=_tournament_timing_preview(tournament_id)
        vn_tz=timezone(timedelta(hours=7))
        now_vn=datetime.now(vn_tz)
        def _future(raw):
            dt=_parse_tournament_dt(raw)
            if not dt: return False
            if dt.tzinfo is None: dt=dt.replace(tzinfo=vn_tz)
            return dt.astimezone(vn_tz)>now_vn
        progress_label=''; progress_deadline=None
        if _future(timing.get('stage1_start_at')):
            progress_label='GĐ1 bắt đầu sau:'; progress_deadline=timing.get('stage1_start_at')
        elif _future(timing.get('stage1_end_at')):
            progress_label='GĐ1 đang diễn ra · thời gian còn lại:'; progress_deadline=timing.get('stage1_end_at')
        elif _future(timing.get('stage1_extension_end_at')):
            progress_label='🔴 GĐ1 đang trong thời gian gia hạn:'; progress_deadline=timing.get('stage1_extension_end_at')
        elif _future(timing.get('league_end_at')):
            progress_label='League Phase · thời gian còn lại:'; progress_deadline=timing.get('league_end_at')
        early_deadline=timing.get('stage1_early_end_at') if _future(timing.get('stage1_early_end_at')) else None

        all_members,_=_safe_rows(db.table('tournament_members').select('user_id').eq('tournament_id',tournament_id).eq('status','active'),'tournament_landing_center_members')
        all_ids=[str(x.get('user_id')) for x in all_members if x.get('user_id')]
        all_users=[]; all_regs=[]
        if all_ids:
            all_users,_=_safe_rows(db.table('users').select('id,username,display_name,is_online,last_seen_at').in_('id',all_ids),'tournament_landing_center_users')
            all_regs,_=_safe_rows(db.table('tournament_registrations').select('user_id,has_host,host_region').eq('tournament_id',tournament_id).in_('user_id',all_ids),'tournament_landing_center_regs')
        all_user_map={str(x.get('id')):x for x in all_users}
        reg_map={str(x.get('user_id')):x for x in all_regs}

        room_rows,_=_safe_rows(db.table('match_rooms').select('*').order('updated_at',desc=True).limit(160),'tournament_landing_center_rooms')
        tournament_match_rows,_=_safe_rows(db.table('tournament_matches').select('*').eq('tournament_id',tournament_id),'tournament_landing_center_match_rows')
        match_map={str(x.get('id')):x for x in tournament_match_rows}
        busy=set(); active_room_statuses={'waiting_ready','playing','friendly_playing','waiting_result_confirm','disputed','confirmed'}
        for r in room_rows:
            if str(r.get('status') or '') in active_room_statuses:
                if r.get('host_user_id'): busy.add(str(r.get('host_user_id')))
                if r.get('guest_user_id'): busy.add(str(r.get('guest_user_id')))
        host_ready=[]
        for mid in all_ids:
            rr=reg_map.get(mid) or {}; uu=all_user_map.get(mid) or {}
            if rr.get('has_host') and is_user_online_now(uu) and mid not in busy:
                host_ready.append({'user_id':mid,'display_name':uu.get('display_name') or uu.get('username') or 'HLV','region':rr.get('host_region') or '—'})

        center_rooms=[]; seen_match_ids=set()
        room_active={'waiting_ready','playing','friendly_playing','waiting_result_confirm','disputed'}
        prefix='TOURNAMENT_ROOM|'
        for r in room_rows:
            if str(r.get('status') or '') not in room_active: continue
            note=str(r.get('note') or '')
            if not note.startswith(prefix): continue
            try: meta=json.loads(note[len(prefix):])
            except Exception: continue
            if str(meta.get('tournament_id') or '')!=str(tournament_id): continue
            mid=str(meta.get('tournament_match_id') or '')
            if not mid or mid in seen_match_ids: continue
            mm=match_map.get(mid) or {}
            if not mm or str(mm.get('status') or '') in {'completed','cancelled'}: continue
            seen_match_ids.add(mid)
            hid=str(r.get('host_user_id') or ''); gid=str(r.get('guest_user_id') or '')
            hn=(all_user_map.get(hid) or {}).get('display_name') or (all_user_map.get(hid) or {}).get('username') or names.get(hid,'HLV')
            gn=(all_user_map.get(gid) or {}).get('display_name') or (all_user_map.get(gid) or {}).get('username') or names.get(gid,'')
            home_name=names.get(str(mm.get('home_user_id') or ''),'HLV')
            away_name=names.get(str(mm.get('away_user_id') or ''),'HLV')
            if gid:
                public_label=f'{hn} vs {gn or away_name}'
                st=str(r.get('status') or '')
                public_status='Đang thi đấu' if st in {'playing','friendly_playing'} else ('Chờ xác nhận kết quả' if st=='waiting_result_confirm' else ('Đang xử lý kết quả' if st=='disputed' else 'Đủ 2 HLV'))
            else:
                expected=away_name if hid==str(mm.get('home_user_id') or '') else home_name
                public_label=f'{hn} · chờ {expected}'; public_status='Chờ đối thủ'
            center_rooms.append({'id':r.get('id'),'status':r.get('status'),'host_user_id':r.get('host_user_id'),'guest_user_id':r.get('guest_user_id'),'host_team':r.get('host_team'),'guest_team':r.get('guest_team'),'public_label':public_label,'public_status':public_status,'tournament_match':mm})

        return {
            'member': member or {'user_id':uid}, 'is_test':is_test, 'days':days, 'mine_set':mine_set,
            'matches':decorated, 'names':names, 'zalos':zalos,
            'stage1_opened':bool(s1_reveals.get(uid)), 'league_mine':league_mine,
            'league_opened':bool(league_reveals.get(uid)), 'test_opponent':test_opponent,
            'test_opponent_days':test_opp_days,
            'progress_label':progress_label,'progress_deadline':progress_deadline,'early_deadline':early_deadline,
            'host_ready':host_ready,'center_rooms':center_rooms,
        }

    @app.get('/tournaments')
    @login_required
    def tournaments():
        opened = tournament_area_enabled()
        tournament_rows = []
        db_ready = True
        if opened:
            tournament_rows, error = _list_tournaments()
            db_ready = not bool(error)
            user = current_user() or {}
            user_id = user.get("id")
            can_admin_manage_tournament = bool(is_admin_user(user))
            for item in tournament_rows:
                item["can_admin_manage"] = can_admin_manage_tournament
                item["my_registration"] = _registration_for_user(item.get("id"), user_id)
                item["my_member"] = _member_for_user(item.get("id"), user_id)
                item["landing_hub"] = _landing_hub_payload(item.get("id"), user, item.get("my_member"))
                item["needs_availability_gate"] = False
                if item.get("my_member") and not can_admin_manage_tournament:
                    av_rows, _ = _safe_rows(
                        db.table("tournament_availability_slots").select("id")
                        .eq("tournament_id", item.get("id")).eq("user_id", user_id),
                        "tournament_landing_availability_gate",
                    )
                    item["needs_availability_gate"] = not bool(av_rows)
                member_rows, _ = _safe_rows(
                    db.table("tournament_members").select("id").eq("tournament_id", item.get("id")).eq("status", "active"),
                    "tournament_member_count",
                )
                item["member_count"] = len(member_rows)
                timing = _tournament_timing_preview(item.get("id"))
                item["stage1_start_at"] = timing.get("stage1_start_at")
                item["stage1_early_end_at"] = timing.get("stage1_early_end_at")
                start_dt = _parse_tournament_dt(item.get("stage1_start_at"))
                if start_dt:
                    if start_dt.tzinfo is None:
                        now_for_start = datetime.now()
                    else:
                        now_for_start = datetime.now(start_dt.tzinfo)
                    item["show_stage1_start_countdown"] = start_dt > now_for_start
                else:
                    item["show_stage1_start_countdown"] = False
                # Public main tournament standings (Stage 1 + League Phase) for the tournament landing page.
                member_rows_full, _ = _safe_rows(
                    db.table("tournament_members").select("*").eq("tournament_id", item.get("id")).eq("status", "active"),
                    "tournament_public_ranking_members",
                )
                member_ids = [str(r.get("user_id")) for r in member_rows_full if r.get("user_id")]
                names = {}
                if member_ids:
                    user_rows, _ = _safe_rows(
                        db.table("users").select("id,username,display_name").in_("id", member_ids),
                        "tournament_public_ranking_users",
                    )
                    names = {str(u.get("id")): (u.get("display_name") or u.get("username") or "HLV") for u in user_rows}

                # Admin quản lý trực tiếp ngay trên BXH Trung tâm.
                # Chỉ nạp dữ liệu quản trị khi tài khoản hiện tại là Admin.
                item["central_admin_players"] = {}
                if can_admin_manage_tournament and member_ids:
                    reg_rows, _ = _safe_rows(
                        db.table("tournament_registrations").select("*")
                        .eq("tournament_id", item.get("id")).in_("user_id", member_ids),
                        "tournament_central_admin_profiles",
                    )
                    reg_map = {str(r.get("user_id")): r for r in reg_rows if r.get("user_id")}
                    member_map = {str(m.get("user_id")): m for m in member_rows_full if m.get("user_id")}
                    admin_match_rows, _ = _safe_rows(
                        db.table("tournament_matches").select("*")
                        .eq("tournament_id", item.get("id")).order("created_at"),
                        "tournament_central_admin_matches",
                    )
                    availability_rows, _ = _safe_rows(
                        db.table("tournament_availability_slots").select("user_id,slot_at")
                        .eq("tournament_id", item.get("id")).in_("user_id", member_ids).order("slot_at"),
                        "tournament_central_admin_availability",
                    )
                    vn_tz = timezone(timedelta(hours=7))
                    now_vn = datetime.now(vn_tz)
                    day_defs = []
                    day_labels = ("Hôm nay", "Ngày mai", "Ngày kia")
                    weekday_names = ("Thứ Hai","Thứ Ba","Thứ Tư","Thứ Năm","Thứ Sáu","Thứ Bảy","Chủ nhật")
                    for offset, label in enumerate(day_labels):
                        d = now_vn.date() + timedelta(days=offset)
                        day_defs.append({
                            "date": d.isoformat(),
                            "label": label,
                            "weekday": f"{weekday_names[d.weekday()]} · {d.strftime('%d/%m')}",
                            "is_weekend": d.weekday() >= 5,
                        })
                    per_player_availability = {uid: {d["date"]: [] for d in day_defs} for uid in member_ids}
                    for av in availability_rows:
                        uid = str(av.get("user_id") or "")
                        if uid not in per_player_availability:
                            continue
                        try:
                            dt = datetime.fromisoformat(str(av.get("slot_at") or "").replace("Z","+00:00"))
                            if dt.tzinfo is None:
                                dt = dt.replace(tzinfo=vn_tz)
                            dt = dt.astimezone(vn_tz)
                        except Exception:
                            continue
                        day_key = dt.date().isoformat()
                        if day_key in per_player_availability[uid] and dt > now_vn:
                            per_player_availability[uid][day_key].append(dt.strftime("%H:%M"))
                    for uid in per_player_availability:
                        for day_key in per_player_availability[uid]:
                            per_player_availability[uid][day_key] = sorted(set(per_player_availability[uid][day_key]))

                    item["central_admin_schedule_matches"] = []
                    for am in admin_match_rows:
                        if not am.get("scheduled_at"):
                            continue
                        item["central_admin_schedule_matches"].append({
                            "id": am.get("id"),
                            "stage_code": am.get("stage_code"),
                            "round_code": am.get("round_code"),
                            "home_name": names.get(str(am.get("home_user_id") or ""), "HLV"),
                            "away_name": names.get(str(am.get("away_user_id") or ""), "HLV"),
                            "scheduled_at": am.get("scheduled_at"),
                            "status": am.get("status"),
                        })

                    per_player_matches = {uid: [] for uid in member_ids}
                    for am in admin_match_rows:
                        h, a = str(am.get("home_user_id") or ""), str(am.get("away_user_id") or "")
                        home_name = names.get(h, "HLV")
                        away_name = names.get(a, "HLV")
                        if h in per_player_matches:
                            row = dict(am)
                            row["home_name"] = home_name
                            row["away_name"] = away_name
                            row["opponent_name"] = away_name
                            row["is_home_for_player"] = True
                            per_player_matches[h].append(row)
                        if a in per_player_matches:
                            row = dict(am)
                            row["home_name"] = home_name
                            row["away_name"] = away_name
                            row["opponent_name"] = home_name
                            row["is_home_for_player"] = False
                            per_player_matches[a].append(row)
                    item["central_admin_schedule_matches"].sort(key=lambda x: str(x.get("scheduled_at") or ""))
                    for uid in member_ids:
                        reg = reg_map.get(uid) or {}
                        mem = member_map.get(uid) or {}
                        item["central_admin_players"][uid] = {
                            "user_id": uid,
                            "display_name": names.get(uid, "HLV"),
                            "registration_id": reg.get("id"),
                            "zalo_name": reg.get("zalo_name") or mem.get("zalo_name") or "",
                            "has_host": bool(reg.get("has_host")),
                            "host_region": reg.get("host_region") or "—",
                            "pot_no": mem.get("pot_no"),
                            "fixed_club_name": mem.get("fixed_club_name") or "",
                            "matches": per_player_matches.get(uid, []),
                            "availability_days": day_defs,
                            "availability_by_day": per_player_availability.get(uid, {}),
                        }
                ranking = {
                    uid: {
                        "user_id": uid,
                        "display_name": names.get(uid, "HLV"),
                        "played": 0,
                        "wins": 0,
                        "draws": 0,
                        "losses": 0,
                        "gf": 0,
                        "ga": 0,
                        "gd": 0,
                        "points": 0,
                        "recent_form": [],
                        "winrate": 0,
                    }
                    for uid in member_ids
                }
                match_rows, _ = _safe_rows(
                    db.table("tournament_matches")
                    .select("home_user_id,away_user_id,home_score,away_score,status,stage_code,completed_at,updated_at,created_at")
                    .eq("tournament_id", item.get("id")).in_("stage_code", ["stage1", "league"]).eq("status", "completed"),
                    "tournament_public_ranking_matches",
                )
                completed_matches = []
                for match in match_rows:
                    h, a = str(match.get("home_user_id")), str(match.get("away_user_id"))
                    if h not in ranking or a not in ranking:
                        continue
                    try:
                        hs, as_ = int(match.get("home_score") or 0), int(match.get("away_score") or 0)
                    except Exception:
                        continue
                    H, A = ranking[h], ranking[a]
                    H["played"] += 1; A["played"] += 1
                    H["gf"] += hs; H["ga"] += as_; A["gf"] += as_; A["ga"] += hs
                    if hs > as_:
                        H["wins"] += 1; H["points"] += 3; A["losses"] += 1
                    elif hs < as_:
                        A["wins"] += 1; A["points"] += 3; H["losses"] += 1
                    else:
                        H["draws"] += 1; A["draws"] += 1; H["points"] += 1; A["points"] += 1
                    completed_matches.append({
                        "home_user_id": h,
                        "away_user_id": a,
                        "home_score": hs,
                        "away_score": as_,
                        "completed_at": match.get("completed_at") or match.get("updated_at") or match.get("created_at"),
                    })

                completed_matches.sort(
                    key=lambda m: _parse_tournament_dt(m.get("completed_at")) or datetime.min,
                    reverse=True,
                )
                for match in completed_matches:
                    for uid, mine, theirs in (
                        (match["home_user_id"], match["home_score"], match["away_score"]),
                        (match["away_user_id"], match["away_score"], match["home_score"]),
                    ):
                        row = ranking.get(uid)
                        if not row or len(row["recent_form"]) >= 5:
                            continue
                        if mine > theirs:
                            pill = {"code": "win", "short": "T", "label": "Thắng"}
                        elif mine < theirs:
                            pill = {"code": "loss", "short": "B", "label": "Bại"}
                        else:
                            pill = {"code": "draw", "short": "H", "label": "Hòa"}
                        row["recent_form"].append(pill)

                public_ranking = list(ranking.values())
                for row in public_ranking:
                    row["gd"] = row["gf"] - row["ga"]
                    total = row["wins"] + row["draws"] + row["losses"]
                    row["winrate"] = round((row["wins"] / total) * 100, 1) if total else 0
                public_ranking.sort(key=lambda r: (r["points"], r["gd"], r["gf"], r["wins"]), reverse=True)
                for idx, row in enumerate(public_ranking, 1):
                    row["rank"] = idx
                item["public_stage1_ranking"] = public_ranking
                status = (item.get("status") or "").lower()
                if status in {"registration", "upcoming"}:
                    item["phase_name"] = "Sắp khai mạc" if item.get("stage1_start_at") else "Đăng ký"
                elif status == "completed":
                    item["phase_name"] = "Đã kết thúc"
                else:
                    item["phase_name"] = "Đang diễn ra"
        return render_template(
            'tournaments.html',
            tournament_open=opened,
            tournaments=tournament_rows,
            tournament_db_ready=db_ready,
        )

    @app.post('/tournaments/<tournament_id>/register')
    @login_required
    def tournament_register(tournament_id):
        if not tournament_area_enabled():
            flash("Khu vực Giải đấu đang tạm đóng.", "warning")
            return redirect(url_for("tournaments"))
        user = current_user() or {}
        user_id = user.get("id")
        tournaments_found, error = _safe_rows(
            db.table("tournaments").select("*").eq("id", tournament_id).limit(1),
            "tournament_register_lookup",
        )
        if error or not tournaments_found:
            flash("Chưa thể truy cập dữ liệu giải đấu. Hãy chạy SQL V1.4.27 trước.", "error")
            return redirect(url_for("tournaments"))
        tournament = tournaments_found[0]
        if not tournament.get("registration_open") or tournament.get("status") not in {"registration", "upcoming"}:
            flash("Giải đấu hiện không nhận đăng ký.", "warning")
            return redirect(url_for("tournaments"))
        if _member_for_user(tournament_id, user_id):
            flash("Bạn đã là thành viên của giải đấu này.", "info")
            return redirect(url_for("tournaments"))
        existing = _registration_for_user(tournament_id, user_id)
        host_choice = (request.form.get("host_choice") or "").strip().lower()
        host_region = (request.form.get("host_region") or "").strip()
        zalo_name = (request.form.get("zalo_name") or "").strip()
        payment_confirmed = request.form.get("payment_confirmed") == "1"
        if host_choice not in {"yes", "no"}:
            flash("Hãy chọn Có Host hoặc Không Host.", "warning")
            return redirect(url_for("tournaments", register=tournament_id))
        if host_region not in {"Bắc", "Trung", "Nam"}:
            flash("Hãy chọn khu vực Bắc, Trung hoặc Nam.", "warning")
            return redirect(url_for("tournaments", register=tournament_id))
        if not zalo_name:
            flash("Hãy nhập Tên Zalo để các HLV có thể liên hệ với bạn.", "warning")
            return redirect(url_for("tournaments", register=tournament_id))
        if len(zalo_name) > 80:
            flash("Tên Zalo tối đa 80 ký tự.", "warning")
            return redirect(url_for("tournaments", register=tournament_id))
        if not payment_confirmed:
            flash("Hãy xác nhận sau khi đã chuyển khoản phí giải.", "warning")
            return redirect(url_for("tournaments", register=tournament_id))
        payload = {
            "tournament_id": tournament_id,
            "user_id": user_id,
            "status": "pending",
            "registered_at": now_iso(),
            "reviewed_at": None,
            "reviewed_by": None,
            "has_host": host_choice == "yes",
            "host_region": host_region,
            "zalo_name": zalo_name,
            "payment_status": "reported",
            "payment_reported_at": now_iso(),
        }
        try:
            if existing:
                execute_query(
                    db.table("tournament_registrations").update(payload).eq("id", existing.get("id")),
                    "tournament_register_update", attempts=2,
                )
            else:
                execute_query(
                    db.table("tournament_registrations").insert(payload),
                    "tournament_register_insert", attempts=2,
                )
            flash("Đã gửi đăng ký. Vui lòng chờ Admin duyệt.", "success")
        except Exception as exc:
            app.logger.exception("Tournament registration failed: %s", exc)
            flash("Không thể gửi đăng ký lúc này.", "error")
        return redirect(url_for("tournaments"))

    @app.post('/tournaments/<tournament_id>/withdraw')
    @login_required
    def tournament_withdraw(tournament_id):
        user = current_user() or {}
        user_id = user.get("id")
        registration = _registration_for_user(tournament_id, user_id)
        if registration and registration.get("status") == "pending":
            execute_query(
                db.table("tournament_registrations").update({"status": "withdrawn"})
                .eq("id", registration.get("id")),
                "tournament_withdraw", attempts=2,
            )
            flash("Đã hủy đăng ký. Bạn có thể đăng ký lại ngay bằng biểu mẫu mới.", "success")
        return redirect(url_for("tournaments"))


    @app.post('/admin/tournaments/design')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_design():
        limits = {
            "hero_cup_width": (80, 420),
            "hero_cup_right": (-120, 260),
            "hero_cup_bottom": (-160, 160),
            "arena_badge_width": (80, 420),
            "arena_badge_x": (0, 100),
            "arena_badge_y": (0, 100),
        }
        current = tournament_design_settings()
        payload = dict(current)
        for key, (low, high) in limits.items():
            try:
                value = int(float(request.form.get(key, payload[key])))
            except (TypeError, ValueError):
                value = payload[key]
            payload[key] = max(low, min(high, value))
        execute_query(
            db.table("system_settings").upsert({
                "setting_key": TOURNAMENT_DESIGN_SETTING_KEY,
                "setting_value": payload,
                "updated_at": now_iso(),
            }, on_conflict="setting_key"),
            "admin_update_tournament_design", attempts=2,
        )
        ttl_cache_delete("tournament_design_settings")
        cache_delete("_tournament_design_settings_cached")
        log_admin_action("Cập nhật bố cục ảnh Giải đấu", "system", details=payload)
        flash("Đã lưu kích thước và vị trí ảnh Giải đấu.", "success")
        return redirect_admin("tournaments")

    @app.post('/tournaments/<tournament_id>/profile')
    @login_required
    def tournament_profile_update(tournament_id):
        user = current_user() or {}
        registration = _registration_for_user(tournament_id, user.get("id"))
        if not registration or registration.get("status") not in {"pending", "approved"}:
            flash("Bạn chưa có đăng ký đang hoạt động trong giải này.", "warning")
            return redirect(url_for("tournaments"))
        host_choice = (request.form.get("host_choice") or "").strip().lower()
        host_region = (request.form.get("host_region") or "").strip()
        zalo_name = (request.form.get("zalo_name") or "").strip()
        if host_choice not in {"yes", "no"} or host_region not in {"Bắc", "Trung", "Nam"} or not zalo_name:
            flash("Hãy nhập đủ Tên Zalo, Khu vực và lựa chọn Host.", "warning")
            return redirect(url_for("tournaments"))
        payload = {"has_host": host_choice == "yes", "host_region": host_region, "zalo_name": zalo_name[:80]}
        try:
            execute_query(
                db.table("tournament_registrations").update(payload).eq("id", registration.get("id")),
                "tournament_profile_update", attempts=2,
            )
        except Exception as exc:
            print(f"tournament_profile_update failed: {exc}")
            flash("Không thể lưu Khu vực / Host / Zalo. Vui lòng báo Admin kiểm tra Database / SQL.", "error")
            return redirect(url_for("tournaments"))
        # Keep member snapshot aligned when present.
        try:
            execute_query(db.table("tournament_members").update({"zalo_name": zalo_name[:80]}).eq("tournament_id", tournament_id).eq("user_id", user.get("id")), "tournament_member_profile_sync", attempts=1)
        except Exception:
            pass
        flash("Đã cập nhật thông tin giải đấu của bạn.", "success")
        return redirect(url_for("tournaments"))

    @app.post('/admin/tournaments/registrations/<registration_id>/profile')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_registration_profile_update(registration_id):
        host_choice = (request.form.get("host_choice") or "").strip().lower()
        host_region = (request.form.get("host_region") or "").strip()
        zalo_name = (request.form.get("zalo_name") or "").strip()
        tournament_id = (request.form.get("tournament_id") or "").strip()
        user_id = (request.form.get("user_id") or "").strip()
        if host_choice not in {"yes", "no"} or host_region not in {"Bắc", "Trung", "Nam"}:
            flash("Thông tin Khu vực/Host không hợp lệ.", "warning")
            return redirect_admin("tournaments")

        reg = None
        reg_id = None if str(registration_id).strip().lower() in {"", "none", "null", "undefined"} else registration_id
        if reg_id:
            rows, _ = _safe_rows(
                db.table("tournament_registrations").select("id,tournament_id,user_id,status").eq("id", reg_id).limit(1),
                "admin_tournament_profile_lookup",
            )
            reg = rows[0] if rows else None

        # HLV added directly by Admin may exist in tournament_members without a registration row.
        # Build a lightweight approved registration profile so Host/region/Zalo can still be edited safely.
        if not reg and tournament_id and user_id:
            rows, _ = _safe_rows(
                db.table("tournament_registrations").select("id,tournament_id,user_id,status").eq("tournament_id", tournament_id).eq("user_id", user_id).limit(1),
                "admin_tournament_profile_lookup_by_member",
            )
            reg = rows[0] if rows else None
            if not reg:
                try:
                    created = execute_query(
                        db.table("tournament_registrations").insert({
                            "tournament_id": tournament_id,
                            "user_id": user_id,
                            "status": "approved",
                            "has_host": host_choice == "yes",
                            "host_region": host_region,
                            "zalo_name": zalo_name[:80] or None,
                            "payment_status": "unreported",
                            "registered_at": now_iso(),
                        }),
                        "admin_tournament_profile_create_for_direct_member", attempts=2,
                    )
                    data = getattr(created, "data", None) or []
                    reg = data[0] if data else None
                except Exception as exc:
                    print(f"admin_tournament_profile_create_for_direct_member failed: {exc}")
                    flash("Không thể tạo hồ sơ thông tin cho HLV. Hãy kiểm tra SQL Database trong Admin.", "error")
                    return redirect_admin("tournaments")

        if not reg:
            flash("Không tìm thấy hồ sơ đăng ký của HLV.", "error")
            return redirect_admin("tournaments")

        reg_id = reg.get("id") or reg_id
        payload = {"has_host": host_choice == "yes", "host_region": host_region, "zalo_name": zalo_name[:80] or None}
        try:
            execute_query(
                db.table("tournament_registrations").update(payload).eq("id", reg_id),
                "admin_tournament_profile_update", attempts=2,
            )
        except Exception as exc:
            print(f"admin_tournament_profile_update failed: {exc}")
            flash("Không thể lưu Khu vực / Host / Zalo. Hãy kiểm tra mục Database / SQL.", "error")
            return redirect_admin("tournaments")

        try:
            execute_query(
                db.table("tournament_members").update({"zalo_name": zalo_name[:80] or None})
                .eq("tournament_id", reg.get("tournament_id")).eq("user_id", reg.get("user_id")),
                "admin_tournament_member_profile_sync", attempts=1,
            )
        except Exception:
            pass
        log_admin_action("Cập nhật thông tin HLV giải đấu", "tournament_registration", target_id=reg_id, details=payload)
        flash("Đã cập nhật Khu vực / Host / Zalo của HLV.", "success")
        return_to = (request.form.get("return_to") or "").strip()
        if return_to == "central" and tournament_id:
            return redirect(url_for("tournaments") + "#ranking")
        if return_to == "tournament" and tournament_id:
            return redirect(url_for('tournaments') + "#bxh")
        return redirect_admin("tournaments")

    @app.get('/admin/tournaments/<tournament_id>/export.xlsx')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_export_excel(tournament_id):
        from io import BytesIO
        from flask import send_file
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment

        regs, _ = _safe_rows(db.table("tournament_registrations").select("*").eq("tournament_id", tournament_id).order("registered_at"), "tournament_export_regs")
        regs = [_decorate_finance(row) for row in _decorate_registration_rows(regs)]
        unmatched, _ = _safe_rows(db.table("tournament_fee_unmatched").select("*").eq("tournament_id", tournament_id).eq("status", "waiting").order("created_at"), "tournament_export_unmatched")
        tours, _ = _safe_rows(db.table("tournaments").select("name").eq("id", tournament_id).limit(1), "tournament_export_name")
        tour_name = (tours[0].get("name") if tours else "PES Arena") or "PES Arena"

        wb = Workbook()
        ws = wb.active
        ws.title = "HLV"
        headers = ["STT","HLV","Trạng thái","Zalo","Khu vực","Host","Đã thu","Phí giải","Trách nhiệm","Thiếu","Cần hoàn thừa","Đã hoàn","Khấu trừ TN","Ghi chú"]
        ws.append(headers)
        for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(horizontal="center")
        active_regs=[r for r in regs if r.get("status") in {"pending","approved"}]
        for i,r in enumerate(active_regs,1):
            paid=int(r.get("amount_paid") or 0)
            fee=min(paid,50000)
            resp=max(0,min(paid-50000,50000))
            ws.append([i,r.get("display_name") or "", "Đã duyệt" if r.get("status")=="approved" else "Chờ duyệt", r.get("zalo_name") or "", r.get("host_region") or "", "Có" if r.get("has_host") else "Không", paid, fee, resp, int(r.get("amount_missing") or 0), int(r.get("amount_surplus") or 0), int(r.get("amount_refunded") or 0), int(r.get("responsibility_deducted") or 0), r.get("payment_note") or ""])
        base_stt = len(active_regs)
        for j,row in enumerate(unmatched,1):
            paid=int(row.get("amount_paid") or 0)
            fee=min(paid,50000)
            resp=max(0,min(paid-50000,50000))
            ws.append([base_stt+j,row.get("payer_name") or "HLV chưa ghép tài khoản","Đã duyệt · Chưa có TK",row.get("zalo_contact") or "","","",paid,fee,resp,max(0,100000-paid),max(0,paid-100000),0,0,row.get("note") or ""])
        for col in "GHIJKLM":
            for cell in ws[col][1:]: cell.number_format = '#,##0" đ"'
        widths=[6,24,14,22,12,10,14,14,14,14,16,14,14,30]
        for idx,w in enumerate(widths,1): ws.column_dimensions[chr(64+idx) if idx<=26 else 'A'].width=w
        ws.freeze_panes="A2"

        wu=wb.create_sheet("Tiền chưa ghép TK")
        uheaders=["STT","Người chuyển / HLV","Zalo / liên hệ","Số tiền","Ghi chú","Trạng thái"]
        wu.append(uheaders)
        for c in wu[1]: c.font=Font(bold=True); c.alignment=Alignment(horizontal="center")
        for i,row in enumerate(unmatched,1): wu.append([i,row.get("payer_name") or "",row.get("zalo_contact") or "",int(row.get("amount_paid") or 0),row.get("note") or "",row.get("status") or "waiting"])
        for cell in wu['D'][1:]: cell.number_format='#,##0" đ"'
        for col,w in zip(['A','B','C','D','E','F'],[6,28,24,16,35,14]): wu.column_dimensions[col].width=w
        wu.freeze_panes="A2"

        ws2=wb.create_sheet("Tổng hợp")
        total=sum(int(r.get("amount_paid") or 0) for r in active_regs) + sum(int(x.get("amount_paid") or 0) for x in unmatched)
        fee_total=sum(min(int(r.get("amount_paid") or 0),50000) for r in active_regs) + sum(min(int(x.get("amount_paid") or 0),50000) for x in unmatched)
        resp_total=sum(max(0,min(int(r.get("amount_paid") or 0)-50000,50000)) for r in active_regs) + sum(max(0,min(int(x.get("amount_paid") or 0)-50000,50000)) for x in unmatched)
        surplus=sum(int(r.get("amount_surplus") or 0) for r in active_regs) + sum(max(0,int(x.get("amount_paid") or 0)-100000) for x in unmatched)
        missing=sum(int(r.get("amount_missing") or 0) for r in active_regs) + sum(max(0,100000-int(x.get("amount_paid") or 0)) for x in unmatched)
        for row in [("Giải",tour_name),("Số HLV theo dõi",len(active_regs)+len(unmatched)),("Trong đó chưa ghép TK",len(unmatched)),("Tổng đã thu",total),("Tiền giải",fee_total),("Tiền trách nhiệm",resp_total),("Cần hoàn thừa",surplus),("Còn thiếu",missing)]: ws2.append(row)
        ws2.column_dimensions['A'].width=28; ws2.column_dimensions['B'].width=28
        for cell in ws2['A']: cell.font=Font(bold=True)
        for cell in ws2['B'][2:]:
            if isinstance(cell.value,(int,float)): cell.number_format='#,##0" đ"'

        output=BytesIO(); wb.save(output); output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"PES_Arena_Le_Phi_Giai_{tournament_id}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    @app.post('/admin/tournaments/registrations/<registration_id>/finance/quick')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_registration_finance_quick(registration_id):
        raw = str(request.form.get("amount_paid") or "0").replace(",", "").replace(".", "").strip()
        try:
            amount = max(0, int(raw))
        except ValueError:
            amount = 0
        if amount not in {0, 50000, 100000, 200000}:
            flash("Mức thu nhanh không hợp lệ.", "warning")
            return redirect_admin("tournaments")
        payload = {
            "amount_paid": amount,
            "fee_amount": 50000,
            "responsibility_amount": 50000,
            "payment_status": "verified" if amount > 0 else "unreported",
            "payment_updated_at": now_iso(),
            "payment_updated_by": (current_user() or {}).get("id"),
        }
        execute_query(db.table("tournament_registrations").update(payload).eq("id", registration_id), "admin_tournament_finance_quick", attempts=2)
        log_admin_action("Cập nhật nhanh lệ phí giải", "tournament_registration", target_id=registration_id, details=payload)
        flash(f"Đã ghi nhận {amount:,}đ.".replace(",", "."), "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/registrations/<registration_id>/finance')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_registration_finance(registration_id):
        def money(name, default=0):
            raw = str(request.form.get(name, default) or "0").replace(",", "").replace(".", "").strip()
            try:
                return max(0, int(raw))
            except ValueError:
                return default
        payload = {
            "amount_paid": money("amount_paid"),
            "fee_amount": money("fee_amount", 50000),
            "responsibility_amount": money("responsibility_amount", 50000),
            "responsibility_deducted": money("responsibility_deducted"),
            "amount_refunded": money("amount_refunded"),
            "payment_note": (request.form.get("payment_note") or "").strip()[:500],
            "payment_updated_at": now_iso(),
            "payment_updated_by": (current_user() or {}).get("id"),
            "payment_status": "verified" if money("amount_paid") > 0 else "unreported",
        }
        execute_query(db.table("tournament_registrations").update(payload).eq("id", registration_id), "admin_tournament_finance_update", attempts=2)
        log_admin_action("Cập nhật lệ phí giải", "tournament_registration", target_id=registration_id, details=payload)
        flash("Đã cập nhật lệ phí HLV.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/<tournament_id>/fees/unmatched/add')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_fee_unmatched_add(tournament_id):
        payer_name = (request.form.get("payer_name") or "").strip()
        zalo_contact = (request.form.get("zalo_contact") or "").strip()
        raw = str(request.form.get("amount_paid") or "0").replace(",", "").replace(".", "").strip()
        try: amount = max(0, int(raw))
        except ValueError: amount = 0
        if not payer_name or amount <= 0:
            flash("Hãy nhập tên người chuyển và số tiền đã nhận.", "warning")
            return redirect_admin("tournaments")
        execute_query(db.table("tournament_fee_unmatched").insert({
            "tournament_id": tournament_id, "payer_name": payer_name, "zalo_contact": zalo_contact or None,
            "amount_paid": amount, "note": (request.form.get("note") or "").strip()[:500],
            "status": "waiting", "created_by": (current_user() or {}).get("id"),
        }), "admin_tournament_fee_unmatched_add", attempts=2)
        flash("Đã lưu khoản tiền chờ ghép tài khoản.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/fees/unmatched/<fee_id>/link')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_fee_unmatched_link(fee_id):
        # V1.4.76: link against ANY player account on the website, including pending accounts.
        user_id = str(request.form.get("user_id") or "").strip()
        fees, _ = _safe_rows(
            db.table("tournament_fee_unmatched").select("*").eq("id", fee_id).limit(1),
            "fee_unmatched_lookup",
        )
        users, _ = _safe_rows(
            db.table("users").select("id,username,display_name,role,account_status").eq("id", user_id).limit(1),
            "fee_link_user_lookup",
        )
        if not fees or not users or users[0].get("role") != "player":
            flash("Không tìm thấy khoản tiền hoặc tài khoản web để ghép.", "error")
            return redirect_admin("tournaments")

        fee = fees[0]
        user = users[0]
        tournament_id = fee.get("tournament_id")
        admin_user = current_user() or {}
        now_value = now_iso()

        regs, _ = _safe_rows(
            db.table("tournament_registrations").select("*")
            .eq("tournament_id", tournament_id).eq("user_id", user_id).limit(1),
            "fee_registration_lookup_by_user",
        )
        reg = regs[0] if regs else None
        if not reg:
            created = execute_query(
                db.table("tournament_registrations").insert({
                    "tournament_id": tournament_id,
                    "user_id": user_id,
                    "status": "approved",
                    "registered_at": now_value,
                    "reviewed_at": now_value,
                    "reviewed_by": admin_user.get("id"),
                    "payment_status": "verified",
                    "amount_paid": int(fee.get("amount_paid") or 0),
                    "payment_updated_at": now_value,
                    "payment_updated_by": admin_user.get("id"),
                }),
                "fee_link_create_registration", attempts=2,
            )
            rows = getattr(created, "data", None) or []
            reg = rows[0] if rows else None
            if not reg:
                regs, _ = _safe_rows(
                    db.table("tournament_registrations").select("*")
                    .eq("tournament_id", tournament_id).eq("user_id", user_id).limit(1),
                    "fee_registration_reload",
                )
                reg = regs[0] if regs else None
        else:
            total = int(reg.get("amount_paid") or 0) + int(fee.get("amount_paid") or 0)
            execute_query(
                db.table("tournament_registrations").update({
                    "status": "approved",
                    "reviewed_at": now_value,
                    "reviewed_by": admin_user.get("id"),
                    "amount_paid": total,
                    "payment_status": "verified",
                    "payment_updated_at": now_value,
                    "payment_updated_by": admin_user.get("id"),
                }).eq("id", reg.get("id")),
                "fee_link_registration", attempts=2,
            )
            reg["amount_paid"] = total
            reg["status"] = "approved"

        if not reg:
            flash("Không thể tạo hồ sơ giải đấu cho tài khoản đã chọn.", "error")
            return redirect_admin("tournaments")

        # Keep the already-approved fee row as an active tournament HLV after linking.
        execute_query(
            db.table("tournament_members").upsert({
                "tournament_id": tournament_id,
                "user_id": user_id,
                "status": "active",
                "approved_at": now_value,
                "approved_by": admin_user.get("id"),
            }, on_conflict="tournament_id,user_id"),
            "fee_link_member", attempts=2,
        )
        execute_query(
            db.table("tournament_fee_unmatched").update({
                "status": "linked",
                "linked_registration_id": reg.get("id"),
                "linked_user_id": user_id,
                "linked_at": now_value,
                "linked_by": admin_user.get("id"),
            }).eq("id", fee_id),
            "fee_link_unmatched", attempts=2,
        )
        label = user.get("display_name") or user.get("username") or "HLV"
        flash(f"Đã ghép khoản tiền vào tài khoản {label}.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/fees/unmatched/<fee_id>/remove')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_fee_unmatched_remove(fee_id):
        # V1.4.80: đây là bản ghi tạm cho HLV đã nộp tiền nhưng CHƯA ghép tài khoản.
        # Admin bấm Xóa nghĩa là bỏ hẳn bản ghi khỏi danh sách theo dõi, không phải
        # chỉ đổi sang cancelled (vì cancelled trước đây vẫn bị query lại và tiếp tục hiện).
        execute_query(
            db.table("tournament_fee_unmatched").delete().eq("id", fee_id),
            "fee_unmatched_delete", attempts=2
        )
        flash("Đã xóa HLV chưa ghép tài khoản khỏi danh sách.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/access')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_access():
        enabled = request.form.get("enabled") == "1"
        execute_query(
            db.table("system_settings").upsert({
                "setting_key": TOURNAMENT_AREA_SETTING_KEY,
                "setting_value": {"enabled": enabled},
                "updated_at": now_iso(),
            }, on_conflict="setting_key"),
            "admin_update_tournament_area", attempts=2,
        )
        ttl_cache_delete("tournament_area_enabled")
        cache_delete("_tournament_area_enabled_cached")
        log_admin_action(
            "Mở khu vực Giải đấu" if enabled else "Tạm đóng khu vực Giải đấu",
            "system", details={"enabled": enabled},
        )
        flash("Đã mở khu vực Giải đấu." if enabled else "Đã tạm đóng khu vực Giải đấu.", "success")
        return redirect_admin("tournaments")

    def _review_registration(registration_id, new_status):
        rows, error = _safe_rows(
            db.table("tournament_registrations").select("*").eq("id", registration_id).limit(1),
            "admin_tournament_registration_lookup",
        )
        if error or not rows:
            flash("Không tìm thấy đăng ký giải đấu.", "error")
            return
        registration = rows[0]
        admin_user = current_user() or {}
        now_value = now_iso()
        execute_query(
            db.table("tournament_registrations").update({
                "status": new_status,
                "reviewed_at": now_value,
                "reviewed_by": admin_user.get("id"),
            }).eq("id", registration_id),
            "admin_tournament_registration_review", attempts=2,
        )
        if new_status == "approved":
            execute_query(
                db.table("tournament_members").upsert({
                    "tournament_id": registration.get("tournament_id"),
                    "user_id": registration.get("user_id"),
                    "status": "active",
                    "approved_at": now_value,
                    "approved_by": admin_user.get("id"),
                    "zalo_name": registration.get("zalo_name"),
                }, on_conflict="tournament_id,user_id"),
                "admin_tournament_member_upsert", attempts=2,
            )
        log_admin_action(
            "Duyệt đăng ký Giải đấu" if new_status == "approved" else "Từ chối đăng ký Giải đấu",
            "tournament_registration",
            details={"registration_id": registration_id, "status": new_status},
        )

    @app.post('/admin/tournaments/registrations/<registration_id>/approve')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_registration_approve(registration_id):
        _review_registration(registration_id, "approved")
        flash("Đã duyệt HLV vào giải đấu.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/registrations/<registration_id>/reject')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_registration_reject(registration_id):
        _review_registration(registration_id, "rejected")
        flash("Đã từ chối đăng ký.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/registrations/<registration_id>/cancel')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_registration_cancel(registration_id):
        rows, error = _safe_rows(
            db.table("tournament_registrations").select("*").eq("id", registration_id).limit(1),
            "admin_tournament_registration_cancel_lookup",
        )
        if error or not rows:
            flash("Không tìm thấy đăng ký giải đấu.", "error")
            return redirect_admin("tournaments")
        registration = rows[0]
        if registration.get("status") == "approved":
            flash("HLV đã được duyệt. Hãy dùng chức năng xóa HLV khỏi giải nếu cần.", "warning")
            return redirect_admin("tournaments")
        execute_query(
            db.table("tournament_registrations").update({
                "status": "withdrawn",
                "reviewed_at": now_iso(),
                "reviewed_by": (current_user() or {}).get("id"),
            }).eq("id", registration_id),
            "admin_tournament_registration_cancel", attempts=2,
        )
        log_admin_action("Hủy đơn đăng ký để HLV đăng ký lại", "tournament_registration", details={"registration_id": registration_id})
        flash("Đã hủy đơn. HLV đã trở về trạng thái Chưa đăng ký và có thể đăng ký lại bằng form mới.", "success")
        return redirect_admin("tournaments")

    @app.post('/admin/tournaments/<tournament_id>/members/add')
    @login_required
    @admin_required
    @admin_permission_required("system_features_manage")
    def admin_tournament_member_add(tournament_id):
        user_id = str(request.form.get("user_id") or "").strip()
        if not user_id:
            flash("Hãy chọn tài khoản cần thêm.", "error")
            return redirect_admin("tournaments")
        admin_user = current_user() or {}
        now_value = now_iso()
        execute_query(
            db.table("tournament_registrations").upsert({
                "tournament_id": tournament_id,
                "user_id": user_id,
                "status": "approved",
                "registered_at": now_value,
                "reviewed_at": now_value,
                "reviewed_by": admin_user.get("id"),
            }, on_conflict="tournament_id,user_id"),
            "admin_tournament_registration_direct", attempts=2,
        )
        execute_query(
            db.table("tournament_members").upsert({
                "tournament_id": tournament_id,
                "user_id": user_id,
                "status": "active",
                "approved_at": now_value,
                "approved_by": admin_user.get("id"),
            }, on_conflict="tournament_id,user_id"),
            "admin_tournament_member_direct", attempts=2,
        )
        log_admin_action("Thêm HLV trực tiếp vào Giải đấu", "tournament_member", details={"tournament_id": tournament_id, "user_id": user_id})
        flash("Đã thêm HLV vào giải đấu.", "success")
        return redirect_admin("tournaments")
